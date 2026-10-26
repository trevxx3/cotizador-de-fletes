#!/usr/bin/env python3
"""
Lee el Excel "COTIZADOR DE FLETES · MAYOREO" y genera los datos que usa index.html
(constante DATOS_EXCEL, sección C). Este script NO contiene datos: todo sale del Excel.

Uso:
  python3 tools/excel_a_json.py RUTA/AL/ARCHIVO.xlsx                       -> imprime el JSON
  python3 tools/excel_a_json.py RUTA/AL/ARCHIVO.xlsx --reemplazar index.html  -> actualiza la página

Requiere:  pip install openpyxl

Qué lee (hoja "Tarifas"):
  D5  Costo por km Camión 3.5 (RAM)      D6  Costo por km Nissan NP300
  D7  Costo por km Tráiler propio        D8  Comidas por viaje
  D9  Permiso municipal por viaje        D10 Sueldo diario chofer (fórmula, se evalúa)
  D11 Tiempo extra por viaje             D12 IVA
  B15:G19  Tipos de llanta: kg, paquetería $/llanta, capacidad Nissan / Camión 3.5 / Tráiler
           (la capacidad de tráiler es un vínculo a la hoja "Capacidades", se toma el valor)
  B23:G..  Destinos: nombre, tipo (Local/Nacional), km sencillo, casetas por viaje,
           fletera sin IVA. (Las columnas E, H, I, J, K son fórmulas y se recalculan en la página.)

IMPORTANTE: el Excel NO se sube al repositorio.
"""
import sys, re, json, ast
import openpyxl

def num(v):
    """Valor numérico de una celda; evalúa fórmulas aritméticas simples como =24000/4/7."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s.startswith("="):
        expr = s[1:]
        if re.fullmatch(r"[0-9.+\-*/() ]+", expr):
            return float(eval(compile(ast.parse(expr, mode="eval"), "f", "eval"), {"__builtins__": {}}))
        return None
    try:
        return float(s.replace(",", "").replace("$", ""))
    except ValueError:
        return None

def limpio(x):
    return None if x is None else (int(x) if float(x).is_integer() else x)

def main():
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    ruta = sys.argv[1]
    wf = openpyxl.load_workbook(ruta)                    # fórmulas
    wv = openpyxl.load_workbook(ruta, data_only=True)    # últimos valores calculados por Excel
    t, tv = wf["Tarifas"], wv["Tarifas"]

    def val(c):
        """Número de la celda: si es fórmula usa el valor calculado; si no hay, la evalúa."""
        f = t[c].value
        if isinstance(f, str) and f.startswith("="):
            v = tv[c].value
            return num(v) if v is not None else num(f)
        return num(f)

    parametros = {
        "costoKmCamion35": limpio(val("D5")),
        "costoKmNissan":   limpio(val("D6")),
        "costoKmTrailer":  limpio(val("D7")),      # vacío en el Excel = sin tarifa
        "comidas":         limpio(val("D8")),
        "permiso":         limpio(val("D9")),
        "sueldoDiario":    limpio(val("D10")),
        "tiempoExtra":     limpio(val("D11")),
        "iva":             limpio(val("D12")),
    }

    tipos = []
    for r in range(15, 20):
        nombre = t[f"B{r}"].value
        if not nombre:
            continue
        tipos.append({
            "nombre": str(nombre).strip(),
            "kg": limpio(val(f"C{r}")),
            "paqueteria": limpio(val(f"D{r}")),
            "capNissan": limpio(val(f"E{r}")),
            "capCamion35": limpio(val(f"F{r}")),
            "capTrailer": limpio(val(f"G{r}")),
        })

    destinos = []
    r = 23
    while r <= t.max_row:
        nombre = t[f"B{r}"].value
        if nombre not in (None, ""):
            destinos.append({
                "nombre": str(nombre).strip(),
                "tipo": str(t[f"C{r}"].value or "").strip(),
                "km": limpio(val(f"D{r}")),
                "casetas": limpio(val(f"F{r}")),
                "fleteraSinIva": limpio(val(f"G{r}")),
            })
        r += 1

    datos = {
        "fuente": "Excel COTIZADOR DE FLETES · MAYOREO (hoja Tarifas)",
        "origen": "Monterrey, N.L.",
        "parametros": parametros,
        "tiposLlanta": tipos,
        "destinos": destinos,
    }
    txt = json.dumps(datos, ensure_ascii=False, separators=(",", ":"))
    if len(sys.argv) >= 4 and sys.argv[2] == "--reemplazar":
        html = open(sys.argv[3], encoding="utf-8").read()
        nuevo, k = re.subn(r"const DATOS_EXCEL = \{.*?\};\s*// FIN DATOS_EXCEL",
                           lambda m: "const DATOS_EXCEL = " + txt + "; // FIN DATOS_EXCEL", html, flags=re.S)
        if k != 1:
            sys.exit("No encontré 'const DATOS_EXCEL = {...}; // FIN DATOS_EXCEL' en " + sys.argv[3])
        open(sys.argv[3], "w", encoding="utf-8").write(nuevo)
        print(f"Actualizado {sys.argv[3]}: {len(tipos)} tipos de llanta, {len(destinos)} destinos")
    else:
        print(json.dumps(datos, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
